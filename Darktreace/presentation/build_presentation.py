from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = Path(__file__).with_name("DarkTrace_Project_Presentation.pptx")
BG = RGBColor(17, 23, 21)
PANEL = RGBColor(27, 37, 32)
PANEL_ALT = RGBColor(31, 47, 39)
INK = RGBColor(234, 241, 235)
MUTED = RGBColor(156, 174, 162)
GREEN = RGBColor(168, 224, 109)
CYAN = RGBColor(112, 214, 199)
AMBER = RGBColor(239, 185, 101)
RED = RGBColor(242, 121, 105)
LINE = RGBColor(61, 79, 68)
FONT = "Aptos"
MONO = "Consolas"


def set_background(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.08))
    band.fill.solid()
    band.fill.fore_color.rgb = GREEN
    band.line.fill.background()


def text_box(slide, text, x, y, w, h, size=18, color=INK, bold=False, font=FONT, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    frame.margin_top = Inches(0.01)
    frame.margin_bottom = Inches(0.01)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def add_card(slide, x, y, w, h, heading, body, accent=GREEN, body_size=15):
    card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    card.fill.solid()
    card.fill.fore_color.rgb = PANEL
    card.line.color.rgb = LINE
    strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.06), Inches(h))
    strip.fill.solid()
    strip.fill.fore_color.rgb = accent
    strip.line.fill.background()
    text_box(slide, heading, x + 0.25, y + 0.2, w - 0.45, 0.34, 16, INK, True)
    text_box(slide, body, x + 0.25, y + 0.68, w - 0.45, h - 0.84, body_size, MUTED)


def add_bullets(slide, items, x=0.85, y=1.75, w=11.5, line_height=0.88, size=19):
    for index, item in enumerate(items):
        top = y + index * line_height
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(top + 0.12), Inches(0.11), Inches(0.11))
        dot.fill.solid()
        dot.fill.fore_color.rgb = GREEN if index % 2 == 0 else CYAN
        dot.line.fill.background()
        text_box(slide, item, x + 0.3, top, w - 0.3, line_height - 0.05, size, INK)


def content_slide(prs, title, eyebrow, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    text_box(slide, eyebrow.upper(), 0.75, 0.42, 10.8, 0.25, 10, GREEN, True, MONO)
    text_box(slide, title, 0.75, 0.78, 11.8, 0.62, 29, INK, True)
    text_box(slide, f"DARKTRACE  /  {number:02d}", 10.65, 7.08, 1.9, 0.18, 8, MUTED, False, MONO, PP_ALIGN.RIGHT)
    return slide


def add_connector(slide, x1, y1, x2, y2, color=GREEN):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(2)
    return line


def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    badge = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.85), Inches(1.05), Inches(0.8), Inches(0.8))
    badge.fill.solid()
    badge.fill.fore_color.rgb = PANEL_ALT
    badge.line.color.rgb = GREEN
    text_box(slide, "DT", 0.85, 1.3, 0.8, 0.28, 18, GREEN, True, MONO, PP_ALIGN.CENTER)
    text_box(slide, "CYBERSECURITY COLLEGE PROJECT  /  2026", 0.85, 2.2, 8.8, 0.3, 11, CYAN, True, MONO)
    text_box(slide, "DarkTrace", 0.82, 2.7, 8.8, 0.95, 45, INK, True)
    text_box(slide, "Simulated Dark-Web Threat Intelligence\n& Security Monitoring Platform", 0.85, 3.8, 9.1, 1.05, 24, GREEN, True)
    text_box(slide, "A safe, local demonstration of detection, risk, and response workflows.", 0.88, 5.2, 8.8, 0.42, 16, MUTED)
    text_box(slide, "FLASK  ·  MYSQL  ·  CHART.JS  ·  THREE.JS", 0.88, 6.65, 8.8, 0.25, 10, MUTED, False, MONO)
    ring = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9.35), Inches(1.65), Inches(2.65), Inches(2.65))
    ring.fill.background()
    ring.line.color.rgb = LINE
    ring.line.width = Pt(1.5)
    inner = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(10.0), Inches(2.3), Inches(1.35), Inches(1.35))
    inner.fill.solid()
    inner.fill.fore_color.rgb = PANEL_ALT
    inner.line.color.rgb = GREEN
    text_box(slide, "LOCAL\nSIMULATION", 10.12, 2.68, 1.1, 0.5, 11, GREEN, True, MONO, PP_ALIGN.CENTER)

    slide = content_slide(prs, "Problem, objectives, and safety", "01 / Project", 2)
    add_card(slide, 0.85, 1.8, 3.75, 3.5, "Problem", "Security teams need a repeatable way to collect observations, prioritize risk, share findings, and track response.", GREEN, 16)
    add_card(slide, 4.8, 1.8, 3.75, 3.5, "Objectives", "Connect the single Admin, MySQL records, risk scores, alerts, response actions, and visual reporting in a testable local app.", CYAN, 15)
    add_card(slide, 8.75, 1.8, 3.75, 3.5, "Safety boundary", "Use fictional signals and account-only controls. No external dark-web access, threat feeds, or network blocking.", AMBER, 16)
    text_box(slide, "ABSTRACT  /  A local simulation of the event-to-alert-to-response lifecycle.", 0.9, 6.15, 11.5, 0.25, 10, MUTED, False, MONO)

    slide = content_slide(prs, "System architecture and data", "02 / Design", 3)
    boxes = [(0.8, "Browser UI", "HTML · CSS\nChart.js · Three.js", CYAN), (3.35, "Flask", "Routes · sessions\nvalidation · CSRF", GREEN), (6.05, "Security rules", "Scoring · alerts\nblocks · audit log", AMBER), (9.0, "MySQL", "SQLAlchemy · PyMySQL\nlocal darktrace DB", RED)]
    for x, heading, body, accent in boxes:
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(2.35), Inches(2.15), Inches(1.55))
        shape.fill.solid()
        shape.fill.fore_color.rgb = PANEL
        shape.line.color.rgb = accent
        text_box(slide, heading, x + 0.15, 2.62, 1.85, 0.32, 16, accent, True, align=PP_ALIGN.CENTER)
        text_box(slide, body, x + 0.15, 3.08, 1.85, 0.55, 12, MUTED, False, align=PP_ALIGN.CENTER)
    add_connector(slide, 2.95, 3.12, 3.28, 3.12)
    add_connector(slide, 5.5, 3.12, 5.98, 3.12)
    add_connector(slide, 8.2, 3.12, 8.93, 3.12)
    text_box(slide, "MySQL stores the Admin, monitored identities, events, alerts, communities, posts, and reports.", 1.1, 4.8, 11.0, 0.35, 15, INK, False, align=PP_ALIGN.CENTER)
    text_box(slide, "Local Chart.js renders metrics · Three.js adds a decorative network · canvas shows entity links", 1.1, 5.35, 11.0, 0.35, 13, CYAN, False, align=PP_ALIGN.CENTER)
    text_box(slide, "No crawler, external feed, or production network control is present.", 1.1, 5.9, 11.0, 0.3, 12, MUTED, False, align=PP_ALIGN.CENTER)

    slide = content_slide(prs, "Risk scoring and response", "03 / Detection", 4)
    steps = [
        (0.85, 1, "Observe", "Synthetic event is recorded", CYAN),
        (3.25, 2, "Detect", "Check login/action rules", GREEN),
        (5.65, 3, "Score", "Bounded 0–100 risk band", AMBER),
        (8.05, 4, "Alert", "Persist event and Admin alert", RED),
        (10.45, 5, "Respond", "Threshold block expires", GREEN),
    ]
    for x, number, heading, body, accent in steps:
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.55), Inches(2.05), Inches(0.52), Inches(0.52))
        circle.fill.solid()
        circle.fill.fore_color.rgb = PANEL_ALT
        circle.line.color.rgb = accent
        text_box(slide, str(number), x + 0.55, 2.2, 0.52, 0.16, 12, accent, True, MONO, PP_ALIGN.CENTER)
        text_box(slide, heading, x, 2.85, 1.65, 0.32, 14, INK, True, align=PP_ALIGN.CENTER)
        text_box(slide, body, x - 0.18, 3.35, 2.0, 0.8, 11, MUTED, False, align=PP_ALIGN.CENTER)
        if number < 5:
            add_connector(slide, x + 1.15, 2.31, x + 2.28, 2.31, LINE)
    text_box(slide, "Risk bands: 0–29 Low · 30–59 Medium · 60–79 High · 80–100 Critical. Threshold and expiry are configurable.", 1.0, 5.4, 11.2, 0.68, 15, GREEN, True, align=PP_ALIGN.CENTER)

    slide = content_slide(prs, "Admin dashboard and controls", "04 / Application", 5)
    add_card(slide, 0.9, 1.85, 5.35, 4.2, "Monitoring", "• Eight security overview metrics\n• Login and alert activity\n• Suspicious-identity review\n• Four Chart.js visualizations\n• User/community/post/report graph", CYAN, 17)
    add_card(slide, 7.05, 1.85, 5.35, 4.2, "Response", "• Force session-generation change\n• Apply or clear temporary blocks\n• Suspend or activate identities\n• Review and acknowledge alerts\n• Update fictional report workflow\n• Audit every admin action", GREEN, 16)
    text_box(slide, "Exactly one Admin signs in; monitored identities have no passwords or login route.", 1.0, 6.45, 11.3, 0.25, 11, MUTED, False, MONO, PP_ALIGN.CENTER)

    slide = content_slide(prs, "Security controls in the prototype", "05 / Safeguards", 6)
    add_bullets(slide, [
        "Admin credentials are environment-configured; passwords are hashed.",
        "HttpOnly/SameSite sessions and session generations protect sign-in.",
        "CSRF tokens and server-side validation protect browser actions.",
        "SQLAlchemy parameter binding and Jinja escaping protect inputs.",
        "Audit rows track block, unblock, suspend, report, and alert actions.",
        "Timed blocks affect local demo records, never IPs or network traffic.",
    ], y=1.8, line_height=0.69, size=16)

    slide = content_slide(prs, "Live demonstration sequence", "06 / Walkthrough", 7)
    add_bullets(slide, [
        "Sign in as the single Admin; review dashboard metrics and charts.",
        "Open the suspicious queue and inspect one fictional identity.",
        "Add a synthetic community event and observe its score and alert.",
        "Review the generated temporary block and countdown.",
        "Unblock or suspend the identity; verify the audit entry.",
        "Update a report and inspect the four-entity relationship graph.",
    ], y=1.75, line_height=0.72, size=16)
    text_box(slide, "Admin login values are read from the private local .env file.", 0.95, 6.4, 11.3, 0.27, 11, AMBER, True, MONO, PP_ALIGN.CENTER)

    slide = content_slide(prs, "Validation, limits, and conclusion", "07 / Evaluation", 8)
    add_card(slide, 0.9, 1.9, 5.35, 3.75, "Verified workflows", "Tests cover single-admin auth, protected routes, seed volumes, risk threshold blocks, expiry/unblock audit, reports, and failed sign-in controls.", GREEN, 16)
    add_card(slide, 7.05, 1.9, 5.35, 3.75, "Limits and takeaway", "Rules can miss context and create false positives. Data is fictional; visual libraries are local. DarkTrace demonstrates defensive workflows, not production monitoring.", AMBER, 16)
    text_box(slide, "Run:  python -m unittest discover -s tests -v", 1.0, 6.35, 11.3, 0.35, 15, CYAN, True, MONO, PP_ALIGN.CENTER)

    prs.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    build_deck()